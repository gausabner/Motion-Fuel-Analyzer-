#!/usr/bin/env python3
"""
run_pipeline.py - Green Basket cc BI pipeline orchestrator (source -> staging -> warehouse).

Steps
  1. (--create-db) create database + login roles          00_create_database.sql
  2. build schemas, staging, warehouse, ETL code, views  01..10 *.sql
  3. EXTRACT: land every raw file as text in stg.*_raw    (COPY, one audit row per file)
  4. TRANSFORM + LOAD: CALL etl.sp_run_full_load(batch)  (05..07)
  5. RECONCILE: CALL etl.sp_reconcile(batch)              (08)
  6. post-load optimisation (indexes, BRIN, aggregation)  (13)
  7. (--with-delta) day-2 incremental master-data batch   (SCD2 versions, inferred members)

Usage (Windows / macOS / Linux, PostgreSQL 14+):
  pip install "psycopg[binary]" pandas openpyxl
  python run_pipeline.py --dsn "postgresql://postgres:<password>@localhost:5432/green_basket_dw" --create-db --with-delta
"""
from __future__ import annotations

import argparse
import csv
import gzip
import io
import json
import re
import sys
import time
from pathlib import Path

import pandas as pd
import psycopg

HERE = Path(__file__).resolve().parent
RAW = HERE.parent / "01_data" / "raw"
SQL_BUILD = ["01_ddl_schemas_audit.sql", "02_ddl_staging_ref.sql", "03_ddl_warehouse.sql", "04_etl_functions.sql",
             "05_etl_transform.sql", "06_etl_load_dimensions.sql", "07_etl_load_facts.sql", "08_reconciliation.sql",
             "09_mart_views.sql", "10_security_grants.sql"]


def log(msg: str):
    print(f"[{time.strftime('%H:%M:%S')}] {msg}", flush=True)


def open_text(path: Path):
    return gzip.open(path, "rt", newline="", encoding="utf-8") if path.suffix == ".gz" else open(path, newline="", encoding="utf-8")


def copy_rows(conn, table: str, cols: list[str], rows) -> int:
    n = 0
    with conn.cursor() as cur:
        with cur.copy(f"COPY {table} ({', '.join(cols)}) FROM STDIN") as cp:
            for r in rows:
                cp.write_row(r)
                n += 1
    return n


def log_extract(conn, batch: int, name: str, target: str, n: int, t0: float):
    with conn.cursor() as cur:
        cur.execute("""INSERT INTO audit.etl_step_log (batch_id, step_name, target_table, rows_in, rows_out, rows_rejected, started_at, message)
                       VALUES (%s, %s, %s, %s, %s, 0, now() - make_interval(secs => %s), 'extract (raw text landing)')""",
                    (batch, f"extract {name}", target, n, n, time.time() - t0))


def none_if_nan(v):
    if v is None:
        return None
    if isinstance(v, float) and pd.isna(v):
        return None
    s = str(v)
    return None if s in ("nan", "NaT", "None") else s


# --------------------------------------------------------------------------- extract
HARVEST_MAP = {
    "A": ["FarmCode", "HarvestDate", "Crop", "Block", "QtyKg", None, "RecordedBy", "Remarks"],
    "B": ["farm", "date", "crop_code", "block_id", "quantity_t", None, "recorder", "notes"],
    "C": ["Farm", "Date", "Product", "Block", "Quantity", "Unit", "Clerk", "Comment"],
}


def extract_harvest(conn, batch):
    for path in sorted((RAW / "internal/farm_logs").glob("harvest_log_*.csv")):
        t0 = time.time()
        with open(path, newline="", encoding="utf-8") as fh:
            first = fh.readline()
            fh.seek(0)
            if first.startswith("FarmCode"):
                tmpl, delim = "A", ","
            elif first.startswith("farm;"):
                tmpl, delim = "B", ";"
            else:
                tmpl, delim = "C", ","
            reader = csv.DictReader(fh, delimiter=delim)
            cols = HARVEST_MAP[tmpl]
            rows = ([path.name, tmpl, i + 1] + [r.get(c) if c else None for c in cols] + [batch]
                    for i, r in enumerate(reader))
            n = copy_rows(conn, "stg.harvest_log_raw",
                          ["source_file", "template_code", "row_num", "farm_txt", "date_txt", "crop_txt", "block_txt",
                           "qty_txt", "unit_txt", "recorder_txt", "remarks_txt", "batch_id"], rows)
        log_extract(conn, batch, path.name, "stg.harvest_log_raw", n, t0)


def extract_csv(conn, batch, path: Path, table: str, cols: list[str], with_row_num=True):
    t0 = time.time()
    with open_text(path) as fh:
        reader = csv.reader(fh)
        header = next(reader)
        idx = [header.index(c) for c in cols]
        rows = ([path.name, i + 1] + [r[j] if r[j] != "" else None for j in idx] + [batch] for i, r in enumerate(reader))
        n = copy_rows(conn, table, ["source_file", "row_num"] + cols + ["batch_id"], rows)
    log_extract(conn, batch, path.name, table, n, t0)
    return n


def extract_master(conn, batch, path: Path, extract_name: str):
    t0 = time.time()
    # keep_default_na=False: pandas would otherwise turn Namibia's ISO code 'NA' into a null (landing must be lossless)
    xl = pd.read_excel(path, sheet_name=None, dtype=str, keep_default_na=False, na_values=[""])
    total = 0

    def land(sheet, table, cols, entity=None, code_col=None):
        nonlocal total
        if sheet not in xl:
            return
        df = xl[sheet]
        rows = []
        for i, r in enumerate(df.itertuples(index=False)):
            d = r._asdict()
            vals = [none_if_nan(d.get(c)) for c in cols]
            if entity:
                rows.append([extract_name, entity, i + 1, none_if_nan(d.get("change_id")), none_if_nan(d.get(code_col))]
                            + vals + [none_if_nan(d.get("approved_by_role")), batch])
            else:
                rows.append([extract_name, i + 1] + vals + [batch])
        if entity:
            n = copy_rows(conn, table, ["extract_name", "entity", "row_num", "change_id", "entity_code"] + cols
                          + ["approved_by_role", "batch_id"], rows)
        else:
            n = copy_rows(conn, table, ["extract_name", "row_num"] + cols + ["batch_id"], rows)
        total += n

    land("Regions", "stg.md_regions_raw", ["region_code", "region_name", "zone_code", "zone_name", "capital_town", "known_aliases"])
    land("Crop_Categories", "stg.md_crop_categories_raw", ["category_code", "category_name", "sector"])
    land("Crops", "stg.md_crops_raw", ["crop_code", "crop_name", "category_code", "scientific_name", "nab_controlled",
                                       "storage_type", "shelf_life_days", "typical_cull_pct"])
    land("Products", "stg.md_products_raw", ["product_code", "product_name", "crop_code", "grade", "pack_type", "sales_uom",
                                             "net_weight_kg", "is_export_grade", "active"])
    land("Farms", "stg.md_farms_raw", ["farm_code", "farm_name", "region", "town", "latitude", "longitude", "elevation_m",
                                       "total_ha", "cultivated_ha", "irrigation_type", "water_source", "farm_manager_code",
                                       "globalgap_certified", "established_year", "harvest_log_template"])
    land("Farm_History", "stg.md_history_raw", ["attribute", "old_value", "new_value", "effective_date"], "FARM", "farm_code")
    land("Customers", "stg.md_customers_raw", ["customer_code", "customer_name", "customer_type", "channel", "town", "region",
                                               "country_code", "currency_code", "credit_terms_days", "key_account",
                                               "onboarding_date", "status"])
    land("Customers_New", "stg.md_customers_raw", ["customer_code", "customer_name", "customer_type", "channel", "town", "region",
                                                   "country_code", "currency_code", "credit_terms_days", "key_account",
                                                   "onboarding_date", "status"])
    land("Customer_History", "stg.md_history_raw", ["attribute", "old_value", "new_value", "effective_date"], "CUSTOMER", "customer_code")
    land("Customer_Merge", "stg.md_customer_merge_raw", ["duplicate_code", "survivor_code", "merge_date", "reason"])
    land("Cost_Categories", "stg.md_cost_categories_raw", ["cost_category_code", "cost_category_name", "cost_group",
                                                           "cost_behaviour", "gl_account", "quantity_uom", "ledger_aliases"])
    log_extract(conn, batch, path.name, "stg.md_*_raw", total, t0)


def extract_weather(conn, batch):
    for path in sorted((RAW / "external/nasa_power").glob("POWER_Point_Daily_*.json")):
        t0 = time.time()
        farm = re.search(r"POWER_Point_Daily_(.+)\.json$", path.name).group(1)
        payload = path.read_text(encoding="utf-8")
        json.loads(payload)  # validate before landing
        with conn.cursor() as cur:
            cur.execute("INSERT INTO stg.weather_json_raw (source_file, farm_code, payload, batch_id) VALUES (%s, %s, %s::jsonb, %s)",
                        (path.name, farm, payload, batch))
        log_extract(conn, batch, path.name, "stg.weather_json_raw", 1, t0)


def extract_pink_sheet(conn, batch):
    from openpyxl import load_workbook
    path = RAW / "external/world_bank/CMO-Historical-Data-Monthly.xlsx"
    t0 = time.time()
    ws = load_workbook(path, read_only=True, data_only=True)["Monthly Prices"]
    rows = list(ws.iter_rows(values_only=True))
    h = next(i for i, r in enumerate(rows) if r and any(str(c).strip() == "Maize" for c in r if c is not None))
    headers, units = rows[h], rows[h + 1]
    out = []
    for r in rows[h + 2:]:
        if not r or r[0] is None or not re.match(r"^\d{4}M\d{2}$", str(r[0])):
            continue
        for j in range(1, len(headers)):
            if headers[j] is None:
                continue
            out.append([path.name, str(r[0]), str(headers[j]), None if units[j] is None else str(units[j]),
                        None if r[j] is None else str(r[j]), batch])
    n = copy_rows(conn, "stg.pink_sheet_raw", ["source_file", "period_txt", "series_header", "unit_txt", "value_txt", "batch_id"], out)
    log_extract(conn, batch, path.name, "stg.pink_sheet_raw", n, t0)


def extract_fred(conn, batch):
    for path in sorted((RAW / "external/fred").glob("*.csv")):
        t0 = time.time()
        with open(path, newline="", encoding="utf-8") as fh:
            reader = csv.reader(fh)
            header = next(reader)
            sid = header[1]
            rows = ([path.name, sid, r[0], r[1] if len(r) > 1 and r[1] != "" else None, batch] for r in reader)
            n = copy_rows(conn, "stg.fred_raw", ["source_file", "series_id", "observation_date", "value_txt", "batch_id"], rows)
        log_extract(conn, batch, path.name, "stg.fred_raw", n, t0)


def extract_all(conn, batch):
    extract_master(conn, batch, RAW / "internal/master_data/GB_Master_Data.xlsx", "main")
    extract_harvest(conn, batch)
    extract_csv(conn, batch, RAW / "internal/farm_logs/planting_register.csv", "stg.planting_register_raw",
                ["cycle_id", "farm_code", "crop_code", "block_id", "planting_date", "planted_ha", "variety",
                 "irrigation_method", "planned_yield_t_ha", "harvested_ha", "status", "last_updated"])
    for p in sorted((RAW / "internal/erp").glob("sales_invoice_lines_FY*.csv.gz")):
        extract_csv(conn, batch, p, "stg.sales_raw",
                    ["invoice_no", "line_no", "doc_type", "invoice_date", "dispatch_date", "customer_code", "product_code",
                     "farm_code", "quantity", "uom", "unit_price", "currency_code", "discount_pct", "unit_std_cost_nad",
                     "freight_nad", "sales_rep_code", "ref_invoice_no"])
    extract_csv(conn, batch, RAW / "internal/stock/stock_movements_monthly.csv", "stg.stock_raw",
                ["farm_code", "product_code", "period", "opening_kg", "received_kg", "dispatched_kg", "shrinkage_kg",
                 "closing_kg", "closing_value_nad", "count_date"])
    extract_csv(conn, batch, RAW / "internal/finance/cost_ledger.csv.gz", "stg.cost_ledger_raw",
                ["posting_id", "posting_date", "farm_code", "cost_centre", "crop_code", "cost_category", "gl_account",
                 "description", "amount_nad", "input_qty", "input_uom", "supplier_code"])
    extract_weather(conn, batch)
    extract_pink_sheet(conn, batch)
    extract_fred(conn, batch)
    extract_csv(conn, batch, RAW / "external/nsa/crop_estimates_by_region.csv", "stg.crop_estimate_raw",
                ["season", "region", "crop", "sector", "area_planted_ha", "production_tonnes", "yield_t_per_ha",
                 "source_table", "release_date"])


# --------------------------------------------------------------------------- main
def run_sql_file(conn, path: Path):
    sql = path.read_text(encoding="utf-8")
    with conn.cursor() as cur:
        cur.execute(sql)


def create_database(dsn: str):
    m = re.match(r"^(postgres(?:ql)?://[^/]+)/([^?]+)", dsn)
    admin_dsn = m.group(1) + "/postgres"
    with psycopg.connect(admin_dsn, autocommit=True) as c:
        exists = c.execute("SELECT 1 FROM pg_database WHERE datname = 'green_basket_dw'").fetchone()
        if not exists:
            c.execute("CREATE DATABASE green_basket_dw ENCODING 'UTF8' TEMPLATE template0")
            log("created database green_basket_dw")
        for role, pw in (("gb_etl_loader", "change-me-etl"), ("gb_pbi_reader", "change-me-reader")):
            if not c.execute("SELECT 1 FROM pg_roles WHERE rolname = %s", (role,)).fetchone():
                c.execute(f"CREATE ROLE {role} LOGIN PASSWORD '{pw}'")
        c.execute("GRANT CONNECT ON DATABASE green_basket_dw TO gb_etl_loader, gb_pbi_reader")


def print_table(conn, sql: str, title: str):
    with conn.cursor() as cur:
        cur.execute(sql)
        df = pd.DataFrame(cur.fetchall(), columns=[d.name for d in cur.description])
    print(f"\n=== {title} ===")
    with pd.option_context("display.max_rows", 200, "display.max_colwidth", 90, "display.width", 250):
        print(df.to_string(index=False))


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--dsn", default="postgresql://postgres:postgres@localhost:5432/green_basket_dw")
    ap.add_argument("--create-db", action="store_true")
    ap.add_argument("--with-delta", action="store_true", help="also run the day-2 incremental master-data batch")
    args = ap.parse_args()

    if args.create_db:
        create_database(args.dsn)
    t_all = time.time()
    with psycopg.connect(args.dsn, autocommit=True) as conn:
        for f in SQL_BUILD:
            log(f"build  {f}")
            run_sql_file(conn, HERE / f)
        batch = conn.execute("SELECT etl.fn_start_batch('FULL', 'Full rebuild from raw extracts (cut-off 2026-09-30)')").fetchone()[0]
        log(f"batch {batch}: extract")
        conn.autocommit = False
        try:
            extract_all(conn, batch)
            conn.commit()
            log(f"batch {batch}: transform + load (CALL etl.sp_run_full_load)")
            conn.execute("CALL etl.sp_run_full_load(%s)", (batch,))
            conn.commit()
            log(f"batch {batch}: reconcile")
            conn.execute("CALL etl.sp_reconcile(%s)", (batch,))
            conn.execute("CALL etl.sp_finish_batch(%s, 'SUCCEEDED')", (batch,))
            conn.commit()
        except Exception:
            conn.rollback()
            conn.execute("CALL etl.sp_finish_batch(%s, 'FAILED')", (batch,))
            conn.commit()
            raise
        conn.autocommit = True
        log("post-load: indexes, BRIN, aggregation table (13_optimisation.sql)")
        run_sql_file(conn, HERE / "13_optimisation.sql")
        conn.execute("ANALYZE")
        print_table(conn, """SELECT step_name, target_table, rows_in, rows_out, rows_rejected, duration_ms, coalesce(message,'') AS message
                             FROM audit.etl_step_log WHERE batch_id = %d AND step_name NOT LIKE 'extract %%' ORDER BY step_id""" % batch,
                    f"Batch {batch} transform/load steps")
        print_table(conn, f"SELECT check_code, check_name, source_value, target_value, status FROM audit.recon_result WHERE batch_id = {batch} ORDER BY recon_id",
                    f"Batch {batch} reconciliation")

        if args.with_delta:
            delta = conn.execute("SELECT etl.fn_start_batch('INCREMENTAL', 'Day-2 master-data delta 2026-10-05')").fetchone()[0]
            conn.autocommit = False
            try:
                extract_master(conn, delta, RAW / "internal/master_data/master_data_delta_2026-10-05.xlsx", "delta")
                conn.execute("CALL etl.sp_run_master_delta(%s)", (delta,))
                conn.execute("CALL etl.sp_reconcile(%s)", (delta,))
                conn.execute("CALL etl.sp_finish_batch(%s, 'SUCCEEDED')", (delta,))
                conn.commit()
            except Exception:
                conn.rollback()
                conn.execute("CALL etl.sp_finish_batch(%s, 'FAILED')", (delta,))
                conn.commit()
                raise
            conn.autocommit = True
            print_table(conn, f"""SELECT step_name, rows_in, rows_out, coalesce(message,'') AS message FROM audit.etl_step_log
                                  WHERE batch_id = {delta} ORDER BY step_id""", f"Batch {delta} (incremental) steps")
    log(f"pipeline finished in {time.time() - t_all:.1f} s")


if __name__ == "__main__":
    sys.exit(main())
