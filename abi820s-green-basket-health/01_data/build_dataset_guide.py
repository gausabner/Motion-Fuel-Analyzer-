#!/usr/bin/env python3
"""
build_dataset_guide.py - makes the project dataset easy to find and to open without any tools.

Writes, from the files themselves (so the counts are always current):
  01_data/README.md                 guide to every dataset file: source, type, rows, columns, real or synthetic
  01_data/GBHS_Dataset_Overview.xlsx an Index sheet, then every source file as a sheet (complete when it has up to
                                    20,000 rows, otherwise its first 1,000 rows), then the clean warehouse tables
                                    exported to 04_powerbi/data_csv (same rule)

Usage: python build_dataset_guide.py
Requires: pandas, openpyxl
"""
from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
RAW = HERE / "raw"
CLEAN = HERE.parent / "04_powerbi" / "data_csv"
FULL_UP_TO, SAMPLE = 20_000, 1_000

# the nine sources, numbered as in the report (section 3.1, data sources table)
SOURCES = [
    ("S1", "PAS ED and clinic attendances", "Transactional", ["internal/pas/pas_ed_opd_encounters_*.csv.gz"],
     "One row per attendance (5 yearly extracts that overlap at year ends)", "Synthetic"),
    ("S2", "PAS inpatient spells", "Transactional", ["internal/pas/pas_inpatient_spells.csv.gz"],
     "Change-data-capture export: one row per spell image (admission, coding, discharge updates)", "Synthetic"),
    ("S3", "PAS patient index and change log", "Transactional / master",
     ["internal/pas/pas_patient_index.csv.gz", "internal/pas/pas_patient_changes.csv.gz"],
     "Pseudonymised patients (no names or IDs) and their residence and payer changes", "Synthetic"),
    ("S4", "Ward midnight census returns", "Transactional (periodic)", ["internal/wards/ward_midnight_census.csv"],
     "One row per ward per day: beds, occupied beds, admissions, discharges", "Synthetic"),
    ("S5", "Nursing e-roster shift summary", "Transactional (periodic)", ["internal/roster/eroster_shift_summary.csv.gz"],
     "One row per care unit per shift: planned and worked hours, overtime, agency, staff on duty", "Synthetic"),
    ("S6", "IDSR weekly returns", "Transactional (periodic)", ["internal/idsr/idsr_weekly_returns.csv"],
     "One row per facility, priority condition and epi week: cases, deaths, tests", "Synthetic"),
    ("S7", "GBHS master data workbook", "Master / reference", ["internal/master_data/GBHS_Master_Data.xlsx"],
     "Regions, facilities, wards and their history, SATS triage scale, ICD-10 codes, IDSR conditions, payers, code mappings",
     "Synthetic (ICD-10 titles from WHO, SATS from the South African Triage Group)"),
    ("S8", "NSA 2023 census, population by region", "External (public)", ["external/nsa/nsa_census2023_population_by_region.csv"],
     "Regional population totals in the published layout (title rows, spaces as thousands separators, footnote)",
     "Real values (Namibia Statistics Agency, 2024)"),
    ("S9", "NASA POWER daily weather", "External (public)", ["external/nasa_power/POWER_Point_Daily_*.json"],
     "Daily rainfall, maximum and minimum temperature and humidity at each facility, in the POWER API layout",
     "Real format; values are stand-ins until fetch_public_data.py is run"),
]


def read_csv(path: Path) -> pd.DataFrame:
    return pd.read_csv(path, dtype=str, keep_default_na=False)


def nasa_long() -> pd.DataFrame:
    """All S9 files in one long table: facility, date and the four weather parameters."""
    frames = []
    for f in sorted(RAW.glob("external/nasa_power/POWER_Point_Daily_*.json")):
        p = json.loads(f.read_text(encoding="utf-8"))["properties"]["parameter"]
        df = pd.DataFrame(p)
        df.insert(0, "date", df.index)
        df.insert(0, "facility_code", f.stem.replace("POWER_Point_Daily_", ""))
        frames.append(df.reset_index(drop=True))
    return pd.concat(frames, ignore_index=True)


def describe(path: Path) -> tuple[int, int, str]:
    """Rows, columns and the column names of one dataset file."""
    if path.suffix == ".json":
        p = json.loads(path.read_text(encoding="utf-8"))["properties"]["parameter"]
        return len(next(iter(p.values()))), len(p) + 1, "date, " + ", ".join(p)
    if path.suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True)
        sheets = [(ws.title, ws.max_row - 1) for ws in wb.worksheets]
        return sum(n for _, n in sheets), len(sheets), "sheets: " + ", ".join(f"{t} ({n})" for t, n in sheets)
    if "nsa_census" in path.name:          # publication layout: title rows before the header
        lines = [l for l in path.read_text(encoding="utf-8").splitlines() if l.strip()]
        data = [l for l in lines if l[:1].isalpha() and "," in l and not l.startswith(("Region", "Source", "Table"))]
        return len(data), 2, "Region, Population (plus title, total and footnote rows)"
    df = read_csv(path)
    return len(df), df.shape[1], ", ".join(df.columns)


def add_sheet(wb: Workbook, name: str, df: pd.DataFrame) -> int:
    """Write df (complete, or its first rows) as a sheet; return the number of rows written."""
    part = df if len(df) <= FULL_UP_TO else df.head(SAMPLE)
    ws = wb.create_sheet(name[:31])
    ws.append(list(part.columns))
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in part.itertuples(index=False):
        ws.append(list(row))
    ws.freeze_panes = "A2"
    for i, col in enumerate(part.columns, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(40, max(10, len(str(col)) + 2))
    return len(part)


def main():
    wb = Workbook()
    index = wb.active
    index.title = "Index"
    idx_rows, guide = [], []
    # ---------------------------------------------------------------- raw sources
    for sid, name, typ, globs, desc, prov in SOURCES:
        files = [f for g in globs for f in sorted(RAW.glob(g))]
        rows_total = 0
        for f in files:
            n, k, cols = describe(f)
            rows_total += n
            rel = f.relative_to(HERE).as_posix()
            sheet = ""
            if f.suffix in (".csv", ".gz") and "nsa_census" not in f.name:
                sheet = f"{sid} {f.name.split('.')[0][-30 + len(sid) + 1:]}"
                written = add_sheet(wb, sheet, read_csv(f))
                idx_rows.append([sheet[:31], sid, name, rel, n, written, prov, desc])
            elif "nsa_census" in f.name:
                sheet = f"{sid} NSA census (as published)"
                lines = pd.DataFrame({"line as published": f.read_text(encoding="utf-8").splitlines()})
                idx_rows.append([sheet[:31], sid, name, rel, n, add_sheet(wb, sheet, lines), prov, desc])
            elif f.suffix == ".xlsx":
                idx_rows.append(["(open the file itself)", sid, name, rel, n, 0, prov, desc])
            if len(files) <= 2:
                guide.append((sid, name, typ, f"`{rel}`", n, k, cols, prov))
        if len(files) > 2:   # one row for a family of files (yearly extracts, one file per facility)
            pattern = globs[0].replace("*", "<year>" if sid == "S1" else "<facility>")
            guide.append((sid, name, typ, f"`raw/{pattern}` ({len(files)} files)", rows_total, k, cols, prov))
        if sid == "S9":
            sheet = "S9 Weather, all facilities"
            df = nasa_long()
            idx_rows.append([sheet, sid, name, "raw/external/nasa_power/*.json (combined)", len(df), add_sheet(wb, sheet, df), prov, desc])
        print(f"{sid} {name}: {len(files)} file(s), {rows_total:,} rows")
    # ---------------------------------------------------------------- day-2 delta
    delta = []
    for f in sorted((RAW / "delta_2026-10-01").glob("*")):
        n, k, cols = describe(f)
        rel = f.relative_to(HERE).as_posix()
        delta.append((rel, n, k))
        if f.suffix == ".csv":
            sheet = "D " + f.stem.replace("_20261001", "")
            idx_rows.append([sheet[:31], "Delta", "Day-2 incremental extract, 1 October 2026", rel, n,
                             add_sheet(wb, sheet, read_csv(f)), "Synthetic", "Rows that arrived after the 30 September cut-off"])
    # ---------------------------------------------------------------- clean warehouse tables
    clean = []
    for f in sorted(CLEAN.glob("*.csv.gz")) if CLEAN.exists() else []:
        df = read_csv(f)
        table = f.name.replace(".csv.gz", "")
        kind = ("Fact" if table.startswith("v_fact") else "Dimension" if table.startswith("v_dim") else
                "Analytics output" if table.startswith("analytics") else "Governance")
        sheet = "W " + table.replace("v_", "")
        written = add_sheet(wb, sheet, df)
        idx_rows.append([sheet[:31], "Warehouse", f"{kind}: {table}", f"04_powerbi/data_csv/{f.name}", len(df), written,
                         "Cleaned and conformed by the ELT", "Analysis-ready table exactly as Power BI imports it"])
        clean.append((table, kind, len(df), df.shape[1]))
    # ---------------------------------------------------------------- index sheet
    index.append(["GBHS dataset overview. Each row is one sheet of this workbook; 'Rows in this sheet' is the complete "
                  f"file when it has up to {FULL_UP_TO:,} rows, otherwise its first {SAMPLE:,} rows. The complete files "
                  "are in 01_data/raw (sources) and 04_powerbi/data_csv (clean warehouse tables)."])
    index.append([])
    index.append(["Sheet", "Source", "Name", "File in the project", "Rows in the file", "Rows in this sheet",
                  "Real or synthetic", "What one row is"])
    for c in index[3]:
        c.font = Font(bold=True)
    for r in idx_rows:
        index.append(r)
    for col, w in zip("ABCDEFGH", (31, 10, 42, 58, 16, 18, 40, 70)):
        index.column_dimensions[col].width = w
    index.freeze_panes = "A4"
    out = HERE / "GBHS_Dataset_Overview.xlsx"
    wb.save(out)
    print(f"wrote {out.name}: {len(wb.sheetnames)} sheets")
    write_readme(guide, delta, clean)


def write_readme(guide, delta, clean):
    man = json.loads((RAW / "generator_manifest.json").read_text(encoding="utf-8"))
    t = []
    t.append("# Dataset: Green Basket Health Services (GBHS)\n")
    t.append("This folder is the project's dataset (brief §4: \"datasets or reproducible links\"). It holds the nine "
             "source extracts that the data warehouse is built from, the day-2 incremental extract, and the files "
             "that document the synthetic data.\n")
    t.append("## Quick start\n")
    t.append(f"- **Open `GBHS_Dataset_Overview.xlsx`** to see the data without any tools. The `Index` sheet lists every "
             f"file; each source is a sheet (complete up to {FULL_UP_TO:,} rows, otherwise its first {SAMPLE:,} rows), "
             "followed by the clean warehouse tables.")
    t.append("- **The complete source files are in `raw/`**, one folder per source system. A `.csv.gz` file is a "
             "compressed CSV: open it with 7-Zip or WinRAR, or read it directly with `pandas.read_csv`.")
    t.append("- **The cleaned, integrated dataset** (star schema, as Power BI imports it) is in "
             "`../04_powerbi/data_csv/` as CSV, and in `../05_backup/` as a PostgreSQL backup.")
    t.append(f"- **Rebuild the dataset** with `python generate_synthetic_data.py`. The seed {man['seed']} makes every "
             "re-run byte-identical.\n")
    t.append("## What is real and what is synthetic\n")
    t.append("- **S1–S7 (internal sources) are synthetic.** GBHS is a fictitious organisation, and no record describes a "
             "real person. Patients carry only a random pseudonym, birth year, sex, residence region and payer class: "
             "there are no names, ID numbers, dates of birth, addresses or clinical notes.")
    t.append("- **S8 holds the real NSA 2023 census totals** by region, in the published layout.")
    t.append("- **S9 has the real NASA POWER file layout**, with stand-in values because the API was not reachable "
             "from the build environment. `python fetch_public_data.py` downloads the real series; no code changes are "
             "needed.")
    t.append("- Data-quality defects were injected on purpose so that the ELT could be tested. `raw/generator_manifest.json` "
             "counts them and `raw/dq_ground_truth.json` lists the affected records (report §3.3 and Appendix E).\n")
    t.append(f"Period: {man['data_start']} to {man['cutoff']}, plus the incremental extract of {man['delta_date']}.\n")
    t.append("## The nine sources\n")
    t.append("| # | Source | Type | File(s) | Rows | Columns | Real or synthetic |")
    t.append("|---|---|---|---|---:|---:|---|")
    for sid, name, typ, rel, n, k, cols, prov in guide:
        t.append(f"| {sid} | {name} | {typ} | {rel} | {n:,} | {k} | {prov} |")
    t.append("\nFor the Excel workbook, rows are the data rows of all its sheets and the column count is its number of "
             "sheets. A NASA POWER row is one day at one facility.\n")
    t.append("## Day-2 incremental extract (1 October 2026)\n")
    t.append("| File | Rows | Columns |")
    t.append("|---|---:|---:|")
    for rel, n, k in delta:
        t.append(f"| `{rel}` | {n:,} | {k} |")
    t.append("\n## Files that document the synthetic data\n")
    t.append("| File | Content |")
    t.append("|---|---|")
    t.append("| `raw/generator_manifest.json` | Seed and software versions, row counts, injected defects per rule, and the "
             "planted behaviour (outbreaks, readmission model, ED wait model, staffing intervention) |")
    t.append("| `raw/dq_ground_truth.json` | The record keys of every injected defect, used to validate the 44 DQ rules |")
    t.append("| `generate_synthetic_data.py` | The generator itself |")
    t.append("| `fetch_public_data.py` | Downloads the real NASA POWER series and checks the NSA totals |")
    t.append("| `build_dataset_guide.py` | Writes this README and the overview workbook |\n")
    if clean:
        t.append("## Clean warehouse tables (`../04_powerbi/data_csv/`)\n")
        t.append("These are the integrated, cleaned tables that the ELT produces and Power BI imports "
                 "(definitions in the report's Appendix A).\n")
        t.append("| Table | Kind | Rows | Columns |")
        t.append("|---|---|---:|---:|")
        for table, kind, n, k in clean:
            t.append(f"| `{table}` | {kind} | {n:,} | {k} |")
        t.append("")
    t.append("## Terms of use\n")
    t.append("The synthetic data may be used freely for teaching and assessment. NSA figures are public statistics, "
             "cited as Namibia Statistics Agency (2024). NASA POWER data are free and open; NASA asks users to "
             "acknowledge the source.")
    (HERE / "README.md").write_text("\n".join(t) + "\n", encoding="utf-8")
    print("wrote README.md")


if __name__ == "__main__":
    main()
